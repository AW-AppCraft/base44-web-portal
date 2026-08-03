#!/usr/bin/env python3
"""
Build SPC_Translator.xlsx - the Excel front end for the translator.

The workbook holds configuration, a run log and instructions. All the work is
done by SPC_Translator_CODE.txt, which is imported into the VBA editor once and
then draws its own buttons.

Cell addresses here must match the Const block at the top of the VBA module.

Run:  python3 build_translator.py [output.xlsx]
"""

import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
NAVY = "1F3864"
BLUE = "2E74B5"
LIGHT = "DCE6F1"

HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor=LIGHT)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")

F_TITLE = Font(name=FONT, size=16, bold=True, color="FFFFFF")
F_SUB = Font(name=FONT, size=10, italic=True, color="404040")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_SECTION = Font(name=FONT, size=11, bold=True, color=NAVY)
F_LABEL = Font(name=FONT, size=10, bold=True)
F_BODY = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_SMALL = Font(name=FONT, size=9)

THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# must match the VBA Const block
MAP_FIRST, MAP_LAST = 25, 74


def banner(ws, title, subtitle, width=8):
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    for c in range(1, width + 1):
        ws.cell(row=1, column=c).fill = HEAD_FILL
    ws.row_dimensions[1].height = 24
    ws["A2"] = subtitle
    ws["A2"].font = F_SUB


def label(ws, row, col, text, font=F_BODY, fill=None, border=None, align=None):
    c = ws.cell(row=row, column=col, value=text)
    c.font = font
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if align:
        c.alignment = align
    return c


def build_config(wb):
    ws = wb.create_sheet("Config")
    banner(ws, "SPC TRANSLATOR - CONFIGURATION",
           "Yellow cells are yours to edit. Grey cells are filled in by the tool.", 4)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 62

    rows = [
        ("Input Folder", r"C:\SPC\inbox",
         "Where the raw gauge / CMM / spreadsheet exports arrive."),
        ("Output Folder", r"C:\SPC\spc-ready",
         "Clean SPC files are written here. Point Settings F8 of the "
         "SPC Calculator at this folder."),
        ("Archive Folder", "",
         "Translated sources are moved here. Blank = <Input Folder>\\archive"),
        ("Rejects Folder", "",
         "Unusable files are moved here. Blank = <Input Folder>\\rejects"),
        ("File Pattern", "*.csv",
         "*.csv, *.xlsx, or *.* to pick up everything."),
        ("Watch Interval (sec)", 30, "Used by Start Watching. Minimum 10."),
        ("Move Source Files", "Y",
         "Y = move sources to the archive folder after translating."),
        ("Max Features", 50, "Most reading columns to take from one row."),
        ("Default Shift", "A", "Used when the source has no shift column."),
        ("Default Operator", "Operator 1", "Used when the source has no operator."),
        ("Default Tool ID", "Tool 1", "Used when the source has no tool / machine."),
        ("Default Date", "",
         "Blank = the date the file is translated. Or fix one as YYYY-MM-DD."),
    ]
    for i, (name, val, note) in enumerate(rows):
        r = 4 + i
        label(ws, r, 1, name, F_LABEL, None, BOX)
        c = ws.cell(row=r, column=2, value=val)
        c.font, c.fill, c.border = F_INPUT, INPUT_FILL, BOX
        label(ws, r, 3, note, F_SMALL, None, BOX,
              Alignment(wrap_text=True, vertical="center"))

    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B10")

    # status block - written by the macro
    label(ws, 17, 1, "Status", F_LABEL, None, BOX)
    for i, (name, val) in enumerate([("Status", "IDLE"), ("Last Run", ""),
                                     ("Files Translated", 0),
                                     ("Files Rejected", 0), ("Rows Written", 0)]):
        r = 17 + i
        label(ws, r, 1, name, F_LABEL, None, BOX)
        c = ws.cell(row=r, column=2, value=val)
        c.font, c.fill, c.border = F_BODY, CALC_FILL, BOX

    # feature map
    label(ws, 23, 1, "FEATURE ORDER (optional)", F_SECTION)
    label(ws, 24, 1, "SPC Feature #", F_HEAD, HEAD_FILL, BOX,
          Alignment(horizontal="center"))
    label(ws, 24, 2, "Source column heading", F_HEAD, HEAD_FILL, BOX)
    label(ws, 24, 3, "Notes", F_HEAD, HEAD_FILL, BOX)
    label(ws, 24, 3,
          "Leave the whole list blank to take reading columns in the order they "
          "appear in the source. Fill it in when files arrive with the columns "
          "in different orders - the name here is matched against the source "
          "heading, and whatever matches lands on that feature number.",
          F_SMALL, HEAD_FILL, BOX, Alignment(wrap_text=True, vertical="top"))
    ws.cell(row=24, column=3).font = Font(name=FONT, size=8, color="FFFFFF")
    ws.row_dimensions[24].height = 46

    for i, r in enumerate(range(MAP_FIRST, MAP_LAST + 1)):
        c = ws.cell(row=r, column=1, value=i + 1)
        c.font, c.border = F_BODY, BOX
        c.alignment = Alignment(horizontal="center")
        c2 = ws.cell(row=r, column=2)
        c2.font, c2.fill, c2.border = F_INPUT, INPUT_FILL, BOX

    ws.freeze_panes = "A4"
    return ws


def build_log(wb):
    ws = wb.create_sheet("Log")
    banner(ws, "TRANSLATION LOG", "One row per file, appended by the macro.", 7)
    heads = ["Translated at", "Source file", "Result", "Rows out",
             "Features", "Output file", "Notes"]
    widths = [20, 30, 11, 10, 34, 30, 70]
    for i, (h, w) in enumerate(zip(heads, widths)):
        ws.column_dimensions[CL(i + 1)].width = w
        label(ws, 3, i + 1, h, F_HEAD, HEAD_FILL, BOX,
              Alignment(horizontal="center", wrap_text=True))
    ws.freeze_panes = "A4"
    return ws


def build_help(wb):
    ws = wb.create_sheet("How To Use")
    banner(ws, "SPC TRANSLATOR - HOW TO USE",
           "Turns any CSV or Excel export into a file the SPC Calculator can "
           "read.", 6)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 76
    ws.column_dimensions["D"].width = 3
    for c in "EFG":
        ws.column_dimensions[c].width = 22

    r = [4]

    def section(text):
        label(ws, r[0], 1, text, F_HEAD, HEAD_FILL, BOX)
        ws.merge_cells(start_row=r[0], start_column=1, end_row=r[0], end_column=3)
        for c in (2, 3):
            ws.cell(row=r[0], column=c).fill = HEAD_FILL
            ws.cell(row=r[0], column=c).border = BOX
        r[0] += 1

    def step(num, head, body):
        if num is not None:
            c = ws.cell(row=r[0], column=1, value=num)
            c.font = Font(name=FONT, size=11, bold=True, color=BLUE)
            c.alignment = Alignment(horizontal="center", vertical="top")
        label(ws, r[0], 2, head, F_LABEL, None, None,
              Alignment(vertical="top", wrap_text=True))
        label(ws, r[0], 3, body, F_BODY, None, None,
              Alignment(vertical="top", wrap_text=True))
        r[0] += 1

    def gap():
        r[0] += 1

    section("SETUP - ONCE")
    step(1, "Save as .xlsm",
         "File > Save As > Excel Macro-Enabled Workbook. Macros cannot be "
         "stored in a .xlsx.")
    step(2, "Import the macro module",
         "Alt+F11 > File > Import File > SPC_Translator.bas. If you were sent "
         "the .txt version, rename it to .bas first, or paste its contents into "
         "a new module and drop the first Attribute line.")
    step(3, "Save, close, reopen",
         "The buttons on the right of this sheet appear automatically. If they "
         "are ever missing, run TR_BuildButtons from Alt+F8.")
    step(4, "Set the folders",
         "Config sheet: B4 = where raw exports arrive, B5 = the clean folder "
         "the SPC Calculator watches. Point Settings F8 of the SPC Calculator "
         "at that same B5 folder.")
    gap()

    section("EVERY DAY")
    step(1, "Preview first",
         "Press Preview. It reports what each file would produce and writes "
         "nothing. Check the feature count matches the number of "
         "characteristics you actually measure.")
    step(2, "Translate",
         "Press Translate Now for a single pass, or Start Watching to keep "
         "checking the input folder every few seconds.")
    step(3, "Check the Log sheet",
         "One row per file: how many rows came out, which columns became which "
         "feature, and every default that was applied.")
    gap()

    section("WHAT IT WORKS OUT BY ITSELF")
    step(None, "Header row",
         "Found automatically, including files with a title and blank lines "
         "above the real table. Files with no header at all work too.")
    step(None, "Separator",
         "Comma, semicolon, tab or pipe - detected per file. Quoted fields are "
         "handled, so a comma inside \"25.03, nominal\" does not break the row.")
    step(None, "Decimal comma",
         "25,0312 is read as 25.0312 when the file uses another separator.")
    step(None, "Which columns are readings",
         "A column counts as a measurement when at least 80% of its values are "
         "numeric and it is not one of the traceability fields. Sample "
         "counters are excluded both by name (No, Index, Sample) and by shape "
         "- a column of 1,2,3... is a counter, not a reading.")
    step(None, "Which columns are traceability",
         "Matched on heading: date / measured / timestamp; shift / turn / team; "
         "operator / op / inspector / user; tool / machine / gauge / equipment "
         "/ station; cluster / batch / lot / serial / part / job.")
    step(None, "Dates",
         "Normalised to YYYY-MM-DD from most common formats and from real "
         "Excel date cells. A trailing time is dropped.")
    gap()

    section("MISSING FIELDS GET DEFAULTS")
    step(None, "Date", "The date the file is translated (Config B15 to fix one)")
    step(None, "Shift", "A  (Config B12)")
    step(None, "Operator", "Operator 1  (Config B13)")
    step(None, "Tool ID", "Tool 1  (Config B14)")
    step(None, "Cluster", "Left blank")
    step(None, "So a bare list of numbers still works",
         "It produces complete rows. Every default applied is named in the Log, "
         "so you always know which traceability is real and which was invented.")
    step(None, "Time is dropped",
         "SPC Data Entry has no time column. If the source has one the Log says "
         "so rather than discarding it silently.")
    gap()

    section("REJECTED FILES")
    step(None, "What gets rejected",
         "Anything that cannot be read, is empty, has no column of readings, or "
         "has no data rows. It is moved to the rejects folder with the reason "
         "in the Log - so nothing malformed ever reaches the SPC Calculator.")
    gap()

    section("KNOWN LIMITS")
    step(None, "Old .xls",
         "Re-save as .xlsx. The file is rejected with that reason rather than "
         "read wrongly.")
    step(None, "One table per file",
         "The first sheet of a workbook, and one table in it. Multi-block "
         "reports need splitting first.")

    ws["E3"] = "BUTTONS"
    ws["E3"].font = F_SECTION
    ws["E4"] = "(appear once macros are enabled)"
    ws["E4"].font = F_SMALL
    return ws


def main(out="SPC_Translator.xlsx"):
    wb = Workbook()
    wb.remove(wb.active)
    build_help(wb)
    build_config(wb)
    build_log(wb)
    wb.active = 0
    wb.save(out)
    print(f"written: {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "SPC_Translator.xlsx")
